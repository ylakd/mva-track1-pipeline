"""End-to-end tests on a tiny synthetic genome (no patient data, no network).

Run:  python -m unittest discover -s tests -v

A synthetic compound-heterozygous pair is planted in a fake gene (FAKEA); the tests check that the
pipeline recovers it, that quality filters remove what they should, and that the NMD rule, gnomAD
parsing, imbalance screen and submission validator behave as specified.
"""
import gzip
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

TMP = tempfile.mkdtemp(prefix="mva_test_")
os.environ["MVA_WORKDIR"] = TMP

from mvapipe import config as cfg  # noqa: E402
from mvapipe import stage01_03_data_qc_coding as s1  # noqa: E402
from mvapipe import stage04_05_annotate_scan as s4  # noqa: E402
from mvapipe import stage06_phenotype as s6  # noqa: E402
from mvapipe import stage07_verify as s7  # noqa: E402
from mvapipe import stage08_locus as s8  # noqa: E402
from mvapipe import stage09_10_nmd_gnomad as s9  # noqa: E402
from mvapipe import stage11_baf as s11  # noqa: E402
from mvapipe import stage12_submission as s12  # noqa: E402
from mvapipe.utils import hgvs_genomic, maxnum  # noqa: E402

GTF_LINES = []


def _gene(chrom, strand, name, gid, exons, mane=True):
    start, end = min(s for s, _ in exons), max(e for _, e in exons)
    attr = f'gene_id "{gid}"; gene_name "{name}"; gene_biotype "protein_coding";'
    tattr = f'{attr} transcript_id "T{gid}"; ' + ('tag "MANE_Select"; ' if mane else "")
    GTF_LINES.append(f"{chrom}\ttest\tgene\t{start}\t{end}\t.\t{strand}\t.\t{attr}")
    GTF_LINES.append(f"{chrom}\ttest\ttranscript\t{start}\t{end}\t.\t{strand}\t.\t{tattr}")
    for s, e in exons:
        for feat in ("exon", "CDS"):
            GTF_LINES.append(f"{chrom}\ttest\t{feat}\t{s}\t{e}\t.\t{strand}\t0\t{tattr}")


_gene("15", "+", "FAKEA", "G1", [(1000, 1200), (2000, 2200), (3000, 3300)])
_gene("15", "+", "FAKEC", "G3", [(10000, 10400)])
_gene("2", "-", "FAKEB", "G2", [(5000, 5200), (6000, 6100), (7000, 7300)])

# (chrom, pos, ref, alt, filter, GT, AD, DP, GQ, in_dbsnp)
VARIANTS = [
    ("2", 7100, "C", "G", "PASS", "0/1", (20, 22), 42, 99, False),     # secondary: stop, ClinVar pathogenic
    ("15", 500, "A", "G", "PASS", "0/1", (20, 20), 40, 99, True),       # non-coding
    ("15", 1100, "T", "G", "PASS", "0/1", (21, 25), 46, 99, False),     # V1: stop (planted)
    ("15", 1150, "G", "A", "PASS", "0/1", (20, 24), 44, 99, True),      # common synonymous
    ("15", 1180, "C", "T", "FS60", "0/1", (20, 20), 40, 99, False),     # filtered -> excluded
    ("15", 1190, "C", "T", "PASS", "0/1", (3, 2), 5, 99, False),        # low depth -> excluded
    ("15", 1195, "C", "T", "PASS", "0/0", (40, 0), 40, 99, False),      # reference -> excluded
    ("15", 1700, "A", "G", "PASS", "0/1", (22, 20), 42, 99, False),     # intronic het between V1 and V2
    ("15", 2150, "T", "G", "PASS", "0/1", (15, 13), 28, 99, False),     # V2: missense (planted)
    ("15", 4000, "G", "T", "PASS", "0/1", (20, 20), 40, 99, False),     # non-coding, outside gene
    ("15", 10100, "A", "C", "PASS", "1/1", (0, 30), 30, 80, False),     # rare homozygous missense
]

MOCK = {
    "chr2:g.7100C>G": {"clinvar": {"rcv": [{"clinical_significance": "Pathogenic"}]},
                       "snpeff": {"ann": [{"effect": "stop_gained", "hgvs_p": "p.Gln33*", "genename": "FAKEB"}]}},
    "chr15:g.1100T>G": {"clinvar": {"rcv": [{"clinical_significance": "Pathogenic"}, {"clinical_significance": "Likely pathogenic"}]},
                        "gnomad_genome": {"af": {"af": 3.3e-5}}, "gnomad_exome": {"af": {"af": 2.8e-5}},
                        "snpeff": {"ann": {"effect": "stop_gained", "hgvs_p": "p.Leu37*", "genename": "FAKEA"}}},
    "chr15:g.1150G>A": {"gnomad_genome": {"af": {"af": 0.30}},
                        "snpeff": {"ann": [{"effect": "synonymous_variant", "hgvs_p": "p.Tyr50Tyr"}]}},
    "chr15:g.2150T>G": {"snpeff": {"ann": [{"effect": "missense_variant", "hgvs_p": "p.Asn717Lys"}, {"effect": "intron_variant"}]},
                        "dbnsfp": {"revel": {"score": [0.47, 0.47]}, "cadd": {"phred": 24.5}}},   # list-valued REVEL (past bug)
    "chr15:g.10100A>C": {"gnomad_exome": {"af": {"af": 1e-5}},
                         "snpeff": {"ann": [{"effect": "missense_variant", "hgvs_p": "p.Thr34Pro"}]}},
}


def write_vcf(dest: Path):
    import pysam
    header = """##fileformat=VCFv4.2
##contig=<ID=2,length=100000>
##contig=<ID=15,length=100000>
##INFO=<ID=DB,Number=0,Type=Flag,Description="dbSNP">
##FILTER=<ID=FS60,Description="FS>60">
##FORMAT=<ID=GT,Number=1,Type=String,Description="GT">
##FORMAT=<ID=AD,Number=R,Type=Integer,Description="AD">
##FORMAT=<ID=DP,Number=1,Type=Integer,Description="DP">
##FORMAT=<ID=GQ,Number=1,Type=Integer,Description="GQ">
#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\tFORMAT\tS1
"""
    order = {"2": 0, "15": 1}
    lines = []
    for ch, pos, ref, alt, flt, gt, ad, dp, gq, db in sorted(VARIANTS, key=lambda v: (order[v[0]], v[1])):
        lines.append(f"{ch}\t{pos}\t.\t{ref}\t{alt}\t500\t{flt}\t{'DB' if db else '.'}\tGT:AD:DP:GQ\t{gt}:{ad[0]},{ad[1]}:{dp}:{gq}")
    plain = dest.with_suffix("")
    plain.write_text(header + "\n".join(lines) + "\n")
    pysam.tabix_compress(str(plain), str(dest), force=True)
    pysam.tabix_index(str(dest), preset="vcf", force=True)
    plain.unlink()


SYN_TERMS = {f"HP:91000{i:02d}": f"Synthetic term {i}" for i in range(1, 9)}


def write_hpo(dest: Path):
    rows = []
    genes = [f"BG{i}" for i in range(300)]
    sizes = {"HP:9100001": 12, "HP:9100002": 60, "HP:9100003": 120, "HP:9100004": 110,
             "HP:9100005": 50, "HP:9100006": 40, "HP:9100007": 30, "HP:9100008": 8}
    rng = np.random.default_rng(1)
    for t in SYN_TERMS:
        for g in rng.choice(genes, sizes[t] - 1, replace=False):
            rows.append((t, "x", 1, g, "OMIM:1"))
    for t in ("HP:9100001", "HP:9100007", "HP:9100003"):
        rows.append((t, "x", 1, "FAKEA", "OMIM:2"))
    rows.append(("HP:9100003", "x", 1, "FAKEB", "OMIM:3"))
    for i in range(30):
        for g in rng.choice(genes, 5 + 7 * i, replace=False):
            rows.append((f"HP:9{i:06d}", "filler", 1, g, "OMIM:4"))
    for g in genes + ["FAKEA", "FAKEB", "FAKEC"]:
        rows.append(("HP:0000118", "Phenotypic abnormality", 1, g, "OMIM:5"))
    pd.DataFrame(rows, columns=["hpo_id", "hpo_name", "ncbi_gene_id", "gene_symbol", "disease_id"]) \
        .to_csv(dest, sep="\t", index=False)


class PipelineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cfg.HPO_TERMS.clear(); cfg.HPO_TERMS.update(SYN_TERMS)
        (cfg.workdir() / "data").mkdir(exist_ok=True)
        write_vcf(cfg.path("vcf"))
        with gzip.open(cfg.path("gtf"), "wt") as f:
            f.write("\n".join(GTF_LINES) + "\n")
        write_hpo(cfg.path("hpo"))

    def test_01_qc_removes_what_it_should(self):
        st = s1.stage02_qc()
        self.assertEqual(st["total_records"], len(VARIANTS))
        self.assertEqual(st["excluded_filter_not_pass"], 1)
        self.assertEqual(st["excluded_reference_or_unknown_gt"], 1)
        self.assertEqual(st["excluded_low_dp_or_gq"], 1)
        self.assertEqual(st["kept"], len(VARIANTS) - 3)
        self.assertEqual(st["hom_alt"], 1)

    def test_02_coding_restriction(self):
        st = s1.stage03_coding()
        self.assertEqual(st["coding_variants"], 5)
        coding = pd.read_parquet(cfg.path("coding"))
        self.assertEqual(sorted(coding["pos"]), [1100, 1150, 2150, 7100, 10100])

    def test_03_annotation_scan_and_ranking(self):
        coding = pd.read_parquet(cfg.path("coding"))
        hg = s4.add_hgvs(coding)
        raw = [dict(MOCK[h], query=h) for h in hg["hgvs"]]
        with open(cfg.path("myvariant_raw"), "w") as fh:
            json.dump(raw, fh)
        s4.stage04_annotate()
        sm = s4.stage05_scan()
        self.assertEqual(sm["gene_assigned_local"], 5)
        self.assertEqual(sm["scan_A_rows"], 1)
        self.assertEqual(sm["scan_B_genes"], 1)
        self.assertEqual(sm["scan_B_genes_with_truncating"], ["FAKEA"])
        res = s6.stage06_phenotype(n_perm=200)
        self.assertEqual(res["top_gene"], "FAKEA")
        self.assertEqual(res["primary"]["gene"], "FAKEA")
        self.assertEqual([v["pos"] for v in res["primary"]["variants"]], [1100, 2150])
        self.assertTrue(res["top_gene_rank_in_all_leave_one_out"])
        with open(cfg.path("selected")) as fh:
            sel = json.load(fh)
        self.assertEqual([s["gene"] for s in sel["secondary"]], ["FAKEB"])
        self.assertTrue(0 <= res["p_any_gene_combined"] <= 1)

    def test_04_parse_handles_list_valued_scores(self):
        out = s4.parse_record(dict(MOCK["chr15:g.2150T>G"], query="chr15:g.2150T>G"))
        self.assertAlmostEqual(out["revel"], 0.47)
        self.assertAlmostEqual(out["cadd"], 24.5)
        self.assertEqual(out["impacto"], 1)
        self.assertTrue(np.isnan(out["gnomad_af"]))

    def test_05_verification_helpers(self):
        c = s7.vcf_call(cfg.path("vcf"), "15", 1100, "T", "G")
        self.assertTrue(c["found"] and c["filter"] == "PASS" and not c["phased"])
        self.assertAlmostEqual(c["alt_fraction"], 25 / 46, places=2)
        self.assertFalse(s7.vcf_call(cfg.path("vcf"), "15", 1100, "T", "A")["found"])
        g = s7.intervening_hets(pd.read_parquet(cfg.path("qc")), "15", 1100, 2150)
        self.assertEqual(g["heterozygous_calls_in_window"], 4)
        self.assertFalse(g["short_read_phasing_feasible"])
        self.assertEqual(s7.residue_from_hgvs_p("p.Asn102Lys"), 102)
        self.assertEqual(s7.residue_from_hgvs_p("p.Leu73*"), 73)

    def test_06_locus_review(self):
        sm = s8.stage08_locus(annotate=False, flank=5000)
        self.assertEqual(sm["gene"], "FAKEA")
        self.assertEqual(sm["records_total"], 9)
        self.assertEqual(sm["by_filter_non_reference"].get("FS60"), 1)
        reg = sm["by_region_non_reference"]
        self.assertEqual(reg.get("intron"), 1)
        self.assertGreaterEqual(reg.get("CDS", 0), 5)

    def test_07_nmd_plus_and_minus_strand(self):
        tx = s9.pick_transcript(cfg.path("gtf"), "FAKEA")
        self.assertEqual(tx["selected_by"], "MANE_Select")
        r1 = s9.nmd_position(tx, 1100)
        self.assertEqual((r1["exon_number"], r1["transcript_offset_nt"]), (1, 100))
        self.assertEqual(r1["distance_to_last_junction_nt"], 302)
        self.assertTrue(r1["predicted_nmd"])
        r2 = s9.nmd_position(tx, 2195)
        self.assertEqual((r2["exon_number"], r2["distance_to_last_junction_nt"]), (2, 6))
        self.assertFalse(r2["predicted_nmd"])
        r3 = s9.nmd_position(tx, 3100)
        self.assertTrue(r3["in_last_exon"])
        self.assertFalse(r3["predicted_nmd"])
        txm = s9.pick_transcript(cfg.path("gtf"), "FAKEB")
        rm = s9.nmd_position(txm, 7100)
        self.assertEqual(rm["exon_number"], 1)
        self.assertEqual(rm["transcript_offset_nt"], 200)
        self.assertEqual(rm["distance_to_last_junction_nt"], 202)
        self.assertTrue(rm["predicted_nmd"])
        self.assertFalse(s9.nmd_position(txm, 6500)["in_exon"])

    def test_08_gnomad_parsing(self):
        class R:
            def __init__(self, code, js): self.status_code, self._js = code, js
            def json(self): return self._js

        class S:
            def __init__(self, resp): self.resp = resp
            def post(self, *a, **k): return self.resp

        present = R(200, {"data": {"variant": {"variant_id": "15-1-A-G",
                          "exome": {"ac": 1, "an": 1000, "homozygote_count": 0}, "genome": None}}})
        r = s9.gnomad_lookup("chr15", 1, "A", "G", session=S(present))
        self.assertEqual(r["status"], "present")
        self.assertAlmostEqual(r["exome"]["af"], 0.001)
        absent = R(200, {"data": {"variant": None}, "errors": [{"message": "Variant not found"}]})
        self.assertEqual(s9.gnomad_lookup("15", 1, "A", "G", session=S(absent))["status"], "absent_from_gnomad_v4")
        self.assertEqual(s9.gnomad_lookup("15", 1, "A", "G", session=S(R(503, {})))["status"], "unavailable")

    def test_09_imbalance_screen_flags_planted_shift(self):
        rng = np.random.default_rng(7)
        rows = []
        for ch in [str(i) for i in range(1, 23)] + ["X"]:
            n = 6000
            pos = np.sort(rng.integers(1, 50_000_000, n))
            dp = rng.poisson(43 if ch != "X" else 22, n).clip(21, 120)
            shift = 0.12 if ch == "21" else 0.0
            af = np.clip(rng.normal(0.5 + shift * rng.choice([-1, 1], n), 0.06, n), 0.05, 0.95)
            ad_alt = np.round(af * dp).astype(int)
            rows.append(pd.DataFrame({"chrom": ch, "pos": pos, "ref": "A", "alt": "G", "qual": 500.0,
                                      "gt": 1, "dp": dp, "gq": 99, "in_dbsnp": True,
                                      "ad_ref": dp - ad_alt, "ad_alt": ad_alt}))
        df = pd.concat(rows, ignore_index=True)
        p = Path(TMP) / "qc_imbalance.parquet"
        df.to_parquet(p)
        sm = s11.stage11_baf(qc_path=p)
        self.assertIn("21", sm["chromosomes_flagged"])
        self.assertNotIn("1", sm["chromosomes_flagged"])
        self.assertGreater(sm["approx_min_detectable_uniform_mosaic_fraction_depth"], 0)
        self.assertLess(sm["x_depth_ratio"], 0.7)

    def test_10_submission_valid_and_validator_catches_errors(self):
        res = s12.stage12_submission()
        self.assertTrue(res["valid"], res["problems"])
        rows = list(pd.read_csv(cfg.path("submission")).itertuples())
        self.assertEqual(rows[0].finding_type, "primary")
        self.assertEqual((rows[0].chrom_1, rows[0].pos_1, rows[0].pos_2), ("chr15", 1100, 2150))
        self.assertGreater(rows[0].epcr, rows[1].epcr)
        text = cfg.path("submission").read_text()
        bad = Path(TMP) / "bad.csv"
        bad.write_text(text.replace("chr15", "15"))
        self.assertTrue(any("chr" in p for p in s12.validate_submission(bad)))
        bad.write_text(text.replace("PROBAND01", "SOMEONE"))
        self.assertTrue(any("proband_id" in p for p in s12.validate_submission(bad)))
        bad.write_text(text.replace(",1100,T,G,", ",1101,T,G,"))
        self.assertTrue(any("not in the VCF" in p for p in s12.validate_submission(bad, cfg.path("vcf"))))

    def test_11_hgvs_and_helpers(self):
        self.assertEqual(hgvs_genomic("17", 7667260, "TA", "T"), "chr17:g.7667261del")
        self.assertEqual(hgvs_genomic("15", 1000, "T", "G"), "chr15:g.1000T>G")
        self.assertEqual(hgvs_genomic("1", 100, "A", "AGT"), "chr1:g.100_101insGT")
        self.assertIsNone(hgvs_genomic("1", 100, "A", "G,T"))
        self.assertTrue(np.isnan(maxnum(None)))
        self.assertEqual(maxnum([[0.1, 0.4], {"score": 0.9}]), 0.9)

    def test_12_hpo_terms_read_from_document(self):
        from docx import Document
        d = Document()
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Synthetic term A", "HP:9100001"
        t.cell(1, 0).text, t.cell(1, 1).text = "Synthetic term C", "HP:9100003"
        p = Path(TMP) / "pheno.docx"
        d.save(p)
        self.assertEqual(s1.extract_hpo_terms(p), {"HP:9100001": "Synthetic term A", "HP:9100003": "Synthetic term C"})


if __name__ == "__main__":
    unittest.main(verbosity=2)
